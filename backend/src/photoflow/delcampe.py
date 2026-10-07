"""Export Delcampe.

Delcampe n'a pas d'API ouverte (API Pass : vendeurs pro avec site marchand). Deux voies :
- Easy Uploader (option Store Plus) : import d'un fichier Excel ou CSV au format impose, colonnes du fichier
  « complet » de Delcampe. Les images sont des URL publiques impossibles a deviner (`/pub/<jeton>.jpg`) que
  Delcampe vient telecharger a l'import.
- Saisie manuelle : `fiche()` renvoie le texte pret a coller.

Voir docs/delcampe.md.
"""

from __future__ import annotations

import csv
import html
import io
import re
import secrets
import uuid
from typing import Literal

from pydantic import BaseModel, Field
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from .models import Caption, Photo


def best_caption(session: Session, photo_id: uuid.UUID) -> Caption | None:
    """La legende humaine prime, sinon la plus recente."""
    caps = session.scalars(select(Caption).where(Caption.photo_id == photo_id)).all()
    if not caps:
        return None
    caps.sort(key=lambda c: (c.source != "human", -(c.updated_at.timestamp() if c.updated_at else 0)))
    return caps[0]


def fiche(session: Session, photo: Photo) -> dict:
    cap = best_caption(session, photo.id)
    data = (cap.data if cap else None) or {}
    lines = []
    if cap and cap.description:
        lines.append(cap.description)
    ep = data.get("epoque") or {}
    lieu = data.get("lieu") or {}
    if ep.get("decennie") and ep.get("decennie") != "inconnue":
        lines.append(f"Epoque estimee : {ep['decennie']} (confiance {ep.get('confiance', '?')}).")
    loc = ", ".join(v for v in [lieu.get("lieu_precis"), lieu.get("ville"), lieu.get("region"), lieu.get("pays")] if v)
    if loc:
        lines.append(f"Lieu suppose : {loc} (confiance {lieu.get('confiance', '?')}).")
    piece = data.get("piece") or {}
    if piece:
        det = ", ".join(str(v) for v in [piece.get("pays"), piece.get("valeur_faciale"), piece.get("type_ou_graveur"), piece.get("annee"), piece.get("metal"), piece.get("atelier")] if v)
        if det:
            lines.append(f"Piece : {det}.")
        if piece.get("etat_estime"):
            lines.append(f"Etat estime : {piece['etat_estime']} (a confirmer en main).")
        if piece.get("annotations"):
            lines.append("Annotations de l'etui : " + " / ".join(piece["annotations"]) + ".")
    cp = data.get("carte_postale") or {}
    if cp:
        det = ", ".join(str(v) for v in [cp.get("editeur"), f"n° {cp['numero']}" if cp.get("numero") else None, "carte-photo" if cp.get("carte_photo") else None] if v)
        if det:
            lines.append(f"Carte postale : {det}.")
        if cp.get("legende_imprimee"):
            lines.append(f"Legende : {cp['legende_imprimee']}")
        if cp.get("voyagee") is not None:
            lines.append("Voyagee" + (f", cachet {cp.get('cachet_lieu') or ''} {cp.get('cachet_date') or ''}".rstrip() if cp.get("voyagee") else "non voyagee") + ".")
    if (data.get("nombre_objets") or 1) > 1 and data.get("lot"):
        lines.append(f"Lot de {data['nombre_objets']} : " + " ; ".join(o.get("resume", "") for o in data["lot"]) + ".")
    if data.get("type_objet") in (None, "photographie"):
        lines.append(f"Format : {photo.width} x {photo.height} px (scan). Tirage original, voir photo pour l'etat.")
    if data.get("etat"):
        lines.append("Defauts : " + ", ".join(data["etat"]) + ".")
    cat_id, _ = category_for(data, photo.nudity_level)
    return {
        "photo_id": str(photo.id),
        "title": (cap.title if cap else None) or photo.filename,
        "description": "\n".join(lines),
        "category": f"{CATEGORY_LABELS[cat_id]} (n° {cat_id})" if cat_id else None,
        "type_objet": data.get("type_objet"),
        "face": data.get("face"),
        "tags": data.get("tags") or [],
        "nudity_level": photo.nudity_level,
        "potentiel": data.get("potentiel"),
        "source": cap.source if cap else None,
    }


# --------------------------------------------------------------------------- categories

# Delcampe > Photographie > Photographies > Photographies (originaux). Numeros releves sur
# https://www.delcampe.net/en_GB/collectables/category-id/photography/photographs/ le 7 oct. 2026. Une categorie
# supprimee reste acceptee 6 mois puis il faut passer a la remplacante : relire la page en cas de rejet.
PHOTO_CATEGORIES: dict[str, int] = {
    "lieux_europe": 34588,
    "lieux_afrique": 12700,
    "lieux_amerique": 34586,
    "lieux_asie": 34585,
    "lieux_oceanie": 34587,
    "lieux": 4220,
    "personnes_anonymes": 22152,
    "personnes_identifiees": 22153,
    "celebrites": 4219,
    "metiers": 18971,
    "militaire": 15593,
    "automobiles": 10956,
    "aviation": 10976,
    "bateaux": 10982,
    "trains": 15558,
    "cyclisme": 15592,
    "sports": 12623,
    "objets": 4221,
    "ethnographie": 9497,
    "pin_up": 15594,
    "avant_1900": 10942,
    "stereoscopie": 12624,
    "autre": 4217,
}

# Nus : rubrique dediee, par epoque (obligatoire sur Delcampe, voir docs/delcampe.md).
NUDE_UNTIL_1920 = 9499
NUDE_1921_1940 = 9500
NUDE_1941_1960 = 9505
NUDE_AFTER_1960 = 30049
NUDE_UNCLASSIFIED = 31175

CATEGORY_LABELS: dict[int, str] = {
    34588: "Photos > Europe",
    12700: "Photos > Afrique",
    34586: "Photos > Amérique",
    34585: "Photos > Asie",
    34587: "Photos > Océanie",
    4220: "Photos > Lieux",
    22152: "Photos > Personnes > Anonymes",
    22153: "Photos > Personnes > Identifiées",
    4216: "Photos > Personnes > Non classées",
    4219: "Photos > Célébrités",
    18971: "Photos > Métiers",
    15593: "Photos > Guerre, militaire",
    10956: "Photos > Automobiles",
    10976: "Photos > Aviation",
    10982: "Photos > Bateaux",
    15558: "Photos > Trains",
    15592: "Photos > Cyclisme",
    12623: "Photos > Sports",
    4221: "Photos > Objets",
    9497: "Photos > Ethniques",
    15594: "Photos > Pin-ups",
    1954: "Photos > Photos signées",
    10942: "Photos > Anciennes (avant 1900)",
    12624: "Photos > Stéréoscopiques",
    4217: "Photos > Autres et non classées",
    NUDE_UNTIL_1920: "Photos > Nus > Nu artistique (…-1920)",
    NUDE_1921_1940: "Photos > Nus > Nu artistique (1921-1940)",
    NUDE_1941_1960: "Photos > Nus > Nu artistique (1941-1960)",
    NUDE_AFTER_1960: "Photos > Nus artistiques (1960-…)",
    NUDE_UNCLASSIFIED: "Photos > Nus > Non classés",
}

_EUROPE = {
    "france", "belgique", "suisse", "luxembourg", "monaco", "italie", "espagne", "portugal", "allemagne",
    "autriche", "pays-bas", "hollande", "royaume-uni", "angleterre", "ecosse", "irlande", "danemark", "suede",
    "norvege", "finlande", "pologne", "grece", "hongrie", "tchecoslovaquie", "roumanie", "yougoslavie", "russie",
}

# Analyses v1 (categorie en texte libre) : rubrique deduite de la scene.
_FROM_SCENE = {
    "portrait": "personnes_anonymes", "groupe": "personnes_anonymes", "militaire": "militaire",
    "travail": "metiers", "sport": "sports",
}


def _decade(data: dict) -> int | None:
    d = ((data.get("epoque") or {}).get("decennie") or "")[:4]
    return int(d) if d.isdigit() else None


def _nude_category(decade: int | None) -> int:
    if decade is None:
        return NUDE_UNCLASSIFIED
    if decade < 1920:
        return NUDE_UNTIL_1920
    if decade < 1940:
        return NUDE_1921_1940
    if decade < 1960:
        return NUDE_1941_1960
    return NUDE_AFTER_1960


def _plain(s: str | None) -> str:
    return (s or "").lower().translate(str.maketrans("éèêëàâäîïôöûüç", "eeeeaaaiioouuc")).strip()


def category_for(data: dict, nudity_level: str | None) -> tuple[int | None, str]:
    """Numero de categorie Delcampe et d'ou il vient ("nudite", "vlm", "scene"). None hors photographies :
    cartes postales, pieces, billets... se rangent dans des arbres trop fins (departement, regne), a choisir a
    l'export."""
    if (data.get("type_objet") or "photographie") != "photographie":
        return None, "a choisir"
    cat = data.get("categorie_delcampe")
    if nudity_level in ("partielle", "integrale") or cat == "nus":
        return _nude_category(_decade(data)), "nudite"
    if cat in PHOTO_CATEGORIES:
        return PHOTO_CATEGORIES[cat], "vlm"
    scene = data.get("scene")
    slug = _FROM_SCENE.get(scene)
    if slug is None and scene in ("paysage", "monument", "scene_de_rue"):
        slug = "lieux_europe" if _plain((data.get("lieu") or {}).get("pays")) in _EUROPE else "lieux"
    if slug is None or slug == "personnes_anonymes":
        dec = _decade(data)
        if dec is not None and dec < 1900:
            slug = "avant_1900"
    return PHOTO_CATEGORIES[slug or "autre"], "scene"


# --------------------------------------------------------------------------- statut de publication

# validee  : fiche relue, prete a partir
# exportee : dans un fichier Easy Uploader (pose a l'export), pas encore confirmee en ligne
# en_vente : en ligne sur Delcampe
# vendue
# retiree  : retiree de la vente ou invendue, peut repartir
STATUSES = ("validee", "exportee", "en_vente", "vendue", "retiree")
# Jamais deux fois en vente : l'export refuse ces photos.
BLOCKING = ("en_vente", "vendue")
# Encore a publier (filtre « a publier » de la grille).
TO_PUBLISH_SQL = "coalesce(p.delcampe_status, '') NOT IN ('exportee', 'en_vente', 'vendue')"


def set_status(session: Session, photo_ids: list[uuid.UUID], status: str | None, only_from: tuple[str | None, ...] | None = None) -> int:
    """Change le statut (None l'efface). only_from : ne touche que les photos dans l'un de ces statuts."""
    if status is not None and status not in STATUSES:
        raise ValueError(f"statut inconnu : {status}")
    cond = ""
    if only_from is not None:
        known = [s for s in only_from if s is not None]
        cond = " AND (delcampe_status = ANY(:from)" + (" OR delcampe_status IS NULL)" if None in only_from else ")")
    at = "NULL" if status is None else "now()"
    params: dict = {"s": status, "ids": photo_ids}
    if only_from is not None:
        params["from"] = known
    r = session.execute(
        text(f"UPDATE photos SET delcampe_status = :s, delcampe_status_at = {at} WHERE id = ANY(:ids){cond}"), params
    )
    return r.rowcount


_REF_IN_FILE = re.compile(r"\bpf-([0-9a-f]{12})\b", re.I)


def refs_in_file(data: bytes, filename: str) -> set[str]:
    """References perso (pf-...) trouvees n'importe ou dans un fichier exporte de Delcampe (Excel ou CSV) :
    on ne depend pas du nom ni de l'ordre des colonnes de leurs exports."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        content = "\n".join(
            str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v is not None
        )
    else:
        try:
            content = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = data.decode("latin-1")
    return {m.lower() for m in _REF_IN_FILE.findall(content)}


def photos_for_refs(session: Session, refs: set[str]) -> dict[str, uuid.UUID]:
    """Reference (12 caracteres hexa) -> photo. Une reference ambigue (deux photos, improbable) est ignoree."""
    if not refs:
        return {}
    rows = session.execute(
        text("SELECT left(replace(id::text, '-', ''), 12) AS ref, id FROM photos "
             "WHERE left(replace(id::text, '-', ''), 12) = ANY(:refs)"),
        {"refs": list(refs)},
    ).all()
    found: dict[str, list[uuid.UUID]] = {}
    for r in rows:
        found.setdefault(r.ref, []).append(r.id)
    return {ref: ids[0] for ref, ids in found.items() if len(ids) == 1}


# --------------------------------------------------------------------------- liens publics

def ensure_public_tokens(session: Session, photos: list[Photo]) -> None:
    """Donne un jeton public aux photos qui n'en ont pas. 32 octets aleatoires : impossible a deviner."""
    for p in photos:
        if not p.public_token:
            p.public_token = secrets.token_urlsafe(32)
    session.flush()


def revoke_public_tokens(session: Session, photo_ids: list[uuid.UUID] | None) -> int:
    """Desactive les liens publics (toutes les photos si photo_ids est None)."""
    if photo_ids is None:
        r = session.execute(text("UPDATE photos SET public_token = NULL WHERE public_token IS NOT NULL"))
    else:
        r = session.execute(
            text("UPDATE photos SET public_token = NULL WHERE public_token IS NOT NULL AND id = ANY(:ids)"),
            {"ids": photo_ids},
        )
    return r.rowcount


def public_url(base_url: str, token: str) -> str:
    return f"{base_url.rstrip('/')}/pub/{token}.jpg"


# --------------------------------------------------------------------------- Easy Uploader

# Colonnes du fichier « complet » Easy Uploader, noms exacts exiges par Delcampe (Easy Uploader - full file
# structure.pdf, centre d'aide Delcampe). En CSV : separateur virgule, texte entre guillemets doubles.
COLUMNS = [
    "category_id", "title", "personal_reference", "description", "selling_type", "price", "minimum_bid_step",
    "initial_quantity", "images", "renew_duration", "renew_total_count", "sale_end_time", "sale_end_day",
    "shipping_model", "weight",
]

TITLE_MAX = 120


class RowOverride(BaseModel):
    price: float | None = Field(default=None, ge=0.01, le=999999)
    category_id: int | None = Field(default=None, ge=1)


class ExportOptions(BaseModel):
    selling_type: Literal["bid", "fixed_price"] = "fixed_price"
    price: float = Field(default=5.0, ge=0.01, le=999999)
    minimum_bid_step: float | None = Field(default=0.5, ge=0.01)
    initial_quantity: int = Field(default=1, ge=1, le=100)
    renew_duration: Literal[7, 10, 14, 21, 28] = 14
    renew_total_count: Literal[0, 1, 2, 3, 4, 5, 10, 99] = 99
    sale_end_time: str | None = Field(default=None, pattern=r"^([01]\d|2[0-3]):[0-5]\d$")
    sale_end_day: int | None = Field(default=None, ge=1, le=7)
    shipping_model: str | None = Field(default=None, max_length=200)
    weight: float | None = Field(default=None, ge=0)


def reference(photo_id: uuid.UUID) -> str:
    """Reference perso Delcampe (20 caracteres max) : retrouve la photo depuis une vente (recherche « pf-... »)."""
    return f"pf-{photo_id.hex[:12]}"


def short_title(title: str) -> str:
    title = re.sub(r"\s+", " ", title).strip()
    if len(title) <= TITLE_MAX:
        return title
    cut = title[: TITLE_MAX - 1].rsplit(" ", 1)[0].rstrip(" ,;:-")
    return cut + "…"


def description_html(description: str) -> str:
    return "".join(f"<p>{html.escape(line.strip(), quote=False)}</p>" for line in description.split("\n") if line.strip())


def preview_row(session: Session, photo: Photo, base_url: str | None = None) -> dict:
    """Une ligne telle qu'elle partira dans le fichier, avec ce qu'il faut verifier avant."""
    f = fiche(session, photo)
    cap = best_caption(session, photo.id)
    data = (cap.data if cap else None) or {}
    cat_id, cat_source = category_for(data, photo.nudity_level)
    warnings = []
    if not cap or not cap.description:
        warnings.append("pas de description")
    if len(f["title"]) > TITLE_MAX:
        warnings.append(f"titre raccourci ({len(f['title'])} > {TITLE_MAX} caracteres)")
    if cat_id is None:
        warnings.append(f"categorie a choisir ({data.get('type_objet')})")
    if data.get("face") == "verso":
        warnings.append("verso")
    if (data.get("nombre_objets") or 1) > 1:
        warnings.append(f"lot de {data['nombre_objets']} objets")
    if photo.delcampe_status == "exportee":
        warnings.append(f"deja exportee le {photo.delcampe_status_at:%d/%m/%Y} : verifier qu'elle n'est pas en ligne")
    return {
        "id": str(photo.id),
        "filename": photo.filename,
        "reference": reference(photo.id),
        "title": short_title(f["title"]),
        "description": f["description"],
        "category_id": cat_id,
        "category_source": cat_source,
        "type_objet": data.get("type_objet"),
        "nudity_level": photo.nudity_level,
        "potentiel": data.get("potentiel"),
        "public_url": public_url(base_url, photo.public_token) if base_url and photo.public_token else None,
        "status": photo.delcampe_status,
        "status_at": photo.delcampe_status_at.isoformat() if photo.delcampe_status_at else None,
        "blocked": photo.delcampe_status in BLOCKING,
        "warnings": warnings,
    }


def build_rows(session: Session, photos: list[Photo], opts: ExportOptions, overrides: dict[uuid.UUID, RowOverride],
               base_url: str) -> tuple[list[list], list[str]]:
    """Lignes du fichier (dans l'ordre de COLUMNS) et ids des photos sans categorie."""
    ensure_public_tokens(session, photos)
    rows, missing = [], []
    for p in photos:
        r = preview_row(session, p)
        o = overrides.get(p.id) or RowOverride()
        cat = o.category_id or r["category_id"]
        if cat is None:
            missing.append(str(p.id))
            continue
        bid = opts.selling_type == "bid"
        rows.append([
            cat,
            r["title"],
            r["reference"],
            description_html(r["description"]),
            opts.selling_type,
            o.price or opts.price,
            (opts.minimum_bid_step or 0.5) if bid else None,
            None if bid else opts.initial_quantity,
            public_url(base_url, p.public_token),
            opts.renew_duration,
            opts.renew_total_count,
            opts.sale_end_time,
            opts.sale_end_day,
            opts.shipping_model or None,
            opts.weight,
        ])
    return rows, missing


def to_csv(rows: list[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=",", quotechar='"', quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    w.writerow(COLUMNS)
    for row in rows:
        w.writerow(["" if v is None else (f"{v:.2f}" if isinstance(v, float) else v) for v in row])
    return buf.getvalue().encode("utf-8")


def to_xlsx(rows: list[list]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Easy Uploader"
    ws.append(COLUMNS)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
